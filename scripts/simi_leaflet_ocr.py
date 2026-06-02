from __future__ import annotations

import argparse
import hashlib
import html
import json
import random
import re
import time
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

LEAFLET_URL = "https://simi.ua/lystivka/"
NETWORK = "СІМ 23 / СІМІ"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) RetailPriceMonitor/0.1"
PRICE_PATTERN = re.compile(r"(?<!\d)\d{1,4}(?:[.,]\d{1,2})?(?!\d)")


@dataclass
class PageReview:
    collection_date: str
    network: str
    page_number: int
    image_url: str
    cache_file: str
    ocr_blocks: int
    status: str


@dataclass
class OcrReview:
    collection_date: str
    network: str
    page_number: int
    image_url: str
    bbox: str
    text: str
    confidence: float
    block_type: str
    manual_review: str


@dataclass
class ReviewError:
    date: str
    network: str
    url: str
    error_type: str
    error_text: str
    manual_review: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download the public Simi leaflet and create a local OCR review workbook."
    )
    parser.add_argument("--limit-pages", type=int, default=0, help="Process only the first N pages. 0 means all pages.")
    parser.add_argument("--skip-ocr", action="store_true", help="Export the image manifest without OCR processing.")
    parser.add_argument("--delay-min", type=float, default=2.0, help="Minimum delay between public HTTP requests.")
    parser.add_argument("--delay-max", type=float, default=5.0, help="Maximum delay between public HTTP requests.")
    parser.add_argument("--retries", type=int, default=3, help="HTTP retry count.")
    return parser.parse_args()


def fetch(session: requests.Session, url: str, args: argparse.Namespace, output: Path | None = None) -> bytes:
    last_error: Exception | None = None
    for attempt in range(1, args.retries + 1):
        try:
            time.sleep(random.uniform(args.delay_min, args.delay_max))
            response = session.get(url, timeout=30)
            response.raise_for_status()
            if output:
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_bytes(response.content)
            return response.content
        except requests.RequestException as error:
            last_error = error
            print(f"[WARN] retry {attempt}/{args.retries} failed for {url}: {error}")
    raise RuntimeError(f"Unable to download {url}: {last_error}")


def extract_current_leaflet_images(page_html: str) -> list[str]:
    soup = BeautifulSoup(page_html, "html.parser")
    slider = soup.select_one(".splide.lystyvka")
    if not slider:
        raise RuntimeError("Current leaflet slider was not found.")
    urls = []
    for image in slider.select(".item-lystyvka img[src]"):
        url = html.unescape(str(image.get("src", "")).strip())
        if url and url not in urls:
            urls.append(url)
    if not urls:
        raise RuntimeError("Current leaflet image pages were not found.")
    return urls


def cache_path(cache_dir: Path, image_url: str) -> Path:
    suffix = Path(urlparse(image_url).path).suffix or ".img"
    digest = hashlib.sha256(image_url.encode("utf-8")).hexdigest()[:20]
    return cache_dir / f"{digest}{suffix}"


def load_ocr_engine() -> Any:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as error:
        raise RuntimeError(
            "rapidocr_onnxruntime is not installed. Run: python -m pip install -r scripts/requirements-simi-ocr.txt"
        ) from error
    return RapidOCR()


def run_ocr(engine: Any, image_file: Path, cache_dir: Path) -> list[list[Any]]:
    cache_file = cache_dir / f"{image_file.stem}.ocr.json"
    if cache_file.exists():
        return json.loads(cache_file.read_text(encoding="utf-8"))
    result, _ = engine(str(image_file))
    blocks = result or []
    cache_file.write_text(json.dumps(blocks, ensure_ascii=False), encoding="utf-8")
    return blocks


def workbook_rows(values: list[Any]) -> list[list[Any]]:
    return [list(asdict(value).values()) for value in values]


def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365F91")
    for column_index, header in enumerate(headers, start=1):
        content_width = max([len(str(header)), *[len(str(row[column_index - 1] or "")) for row in rows]], default=12)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(content_width + 2, 12), 60)


def write_workbook(
    exports_dir: Path,
    pages: list[PageReview],
    ocr_blocks: list[OcrReview],
    errors: list[ReviewError],
    leaflet_url: str,
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(
        workbook,
        "Leaflet Pages",
        ["Collection Date", "Network", "Page Number", "Image URL", "Cache File", "OCR Blocks", "Status"],
        workbook_rows(pages),
    )
    add_sheet(
        workbook,
        "OCR Review",
        ["Collection Date", "Network", "Page Number", "Image URL", "BBox", "OCR Text", "Confidence", "Block Type", "Manual Review"],
        workbook_rows(ocr_blocks),
    )
    add_sheet(
        workbook,
        "Errors",
        ["Date", "Network", "URL", "Error Type", "Error Text", "Manual Review"],
        workbook_rows(errors),
    )
    add_sheet(
        workbook,
        "Summary",
        ["Metric", "Value"],
        [
            ["Leaflet URL", leaflet_url],
            ["Page Count", len(pages)],
            ["OCR Block Count", len(ocr_blocks)],
            ["Price Candidate Block Count", sum(block.block_type == "PRICE_CANDIDATE" for block in ocr_blocks)],
            ["Error Count", len(errors)],
            ["Source Status", "manual review required"],
        ],
    )
    exports_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%dT%H-%M-%S")
    output = exports_dir / f"simi-leaflet-review-{timestamp}.xlsx"
    workbook.save(output)
    return output


def main() -> int:
    args = parse_args()
    if args.delay_min < 0 or args.delay_max < args.delay_min:
        raise ValueError("Invalid delay range.")
    collection_date = datetime.now().date().isoformat()
    cache_dir = Path(".cache") / "simi"
    exports_dir = Path("exports")
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    errors: list[ReviewError] = []
    pages: list[PageReview] = []
    reviews: list[OcrReview] = []

    page_html = fetch(session, LEAFLET_URL, args).decode("utf-8", errors="replace")
    image_urls = extract_current_leaflet_images(page_html)
    if args.limit_pages > 0:
        image_urls = image_urls[: args.limit_pages]
    print(f"[INFO] found {len(image_urls)} current Simi leaflet pages")

    engine = None if args.skip_ocr else load_ocr_engine()
    for page_number, image_url in enumerate(image_urls, start=1):
        image_file = cache_path(cache_dir, image_url)
        try:
            if not image_file.exists():
                fetch(session, image_url, args, image_file)
            blocks = [] if engine is None else run_ocr(engine, image_file, cache_dir)
            pages.append(
                PageReview(collection_date, NETWORK, page_number, image_url, str(image_file), len(blocks), "manual review required")
            )
            for bbox, text, confidence in blocks:
                clean_text = str(text).strip()
                reviews.append(
                    OcrReview(
                        collection_date,
                        NETWORK,
                        page_number,
                        image_url,
                        json.dumps(bbox, ensure_ascii=False),
                        clean_text,
                        round(float(confidence), 4),
                        "PRICE_CANDIDATE" if PRICE_PATTERN.search(clean_text) else "OCR_TEXT",
                        "yes",
                    )
                )
            print(f"[INFO] page {page_number}/{len(image_urls)}: {len(blocks)} OCR blocks")
        except Exception as error:  # Continue processing the remaining public pages.
            errors.append(
                ReviewError(
                    datetime.now(timezone.utc).isoformat(),
                    NETWORK,
                    image_url,
                    "LEAFLET_PAGE_ERROR",
                    str(error),
                    "yes",
                )
            )
            print(f"[ERROR] page {page_number}/{len(image_urls)}: {error}")

    output = write_workbook(exports_dir, pages, reviews, errors, LEAFLET_URL)
    print(f"[INFO] OCR review export created: {output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
